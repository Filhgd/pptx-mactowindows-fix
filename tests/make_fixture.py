#!/usr/bin/env python3
"""Builds a test presentation that looks like one saved by PowerPoint for Mac:
a pasted PDF clip stored as an EMF with the PDF inside a GDIC comment.

Usage: make_fixture.py OUT_DIR
Writes: mac_clip.pptx (needs fixing) and plain.pptx (nothing to fix)."""
import io, os, struct, sys, zipfile
from pptx import Presentation
from pptx.util import Inches
from PIL import Image, ImageDraw

PDF_W, PDF_H = 400, 270


def make_pdf():
    content = (b"q 0 0 0 RG 4 w 20 20 m 380 250 l S Q "
               b"BT /F1 40 Tf 30 120 Td (SHARP 123) Tj ET "
               b"0 0 1 rg 300 30 60 60 re f")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 %d %d] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>" % (PDF_W, PDF_H),
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for o in offsets:
        out += b"%010d 00000 n \n" % o
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return out


def make_emf(pdf):
    # EMR_HEADER (108 bytes), same layout as PowerPoint for Mac writes.
    header = bytearray(bytes.fromhex(
        "010000006c00000000000000000000008f0100000d01000000000000000000000c1f0000ee140000"
        "20454d4600000100b8df07000300000001000000000000000000000000000000e8050000d6030000"
        "2d010000c3000000000000000000000000000000c8970400b8f90200"))
    # EMR_COMMENT_MULTIFORMATS ("GDIC") with one PDF format.
    body = (b"GDIC" + struct.pack("<I", 0x40000004) + b"\0" * 16 + struct.pack("<I", 1)
            + b" FDP" + struct.pack("<III", 0, len(pdf), 44) + pdf)
    padded = body + b"\0" * (-len(body) % 4)
    comment = struct.pack("<III", 70, 12 + len(padded), len(body)) + padded
    eof = bytes.fromhex("0e000000100000000000000000000000")
    total = len(header) + len(comment) + len(eof)
    struct.pack_into("<II", header, 48, total, 3)
    return bytes(header) + comment + eof


def png_bytes(size, color):
    im = Image.new("RGB", size, color)
    ImageDraw.Draw(im).text((10, 10), "fallback", fill=(0, 0, 0))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def build(out_dir):
    os.makedirs(out_dir, exist_ok=True)
    clip_png = png_bytes((PDF_W, PDF_H), (200, 200, 200))
    other_png = png_bytes((300, 200), (250, 220, 180))

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    s1.shapes.add_picture(io.BytesIO(clip_png), Inches(1), Inches(1), width=Inches(8))
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    pic = s2.shapes.add_picture(io.BytesIO(clip_png), Inches(1), Inches(1), width=Inches(4))
    pic.crop_left = 0.2
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    s3.shapes.add_picture(io.BytesIO(other_png), Inches(1), Inches(1), width=Inches(4))
    plain = os.path.join(out_dir, "plain.pptx")
    prs.save(plain)

    # Turn the clip picture into a Mac-style EMF.
    zin = zipfile.ZipFile(plain)
    clip_part = next(n for n in zin.namelist() if n.startswith("ppt/media/") and zin.read(n) == clip_png)
    emf_part = clip_part.rsplit(".", 1)[0] + ".emf"
    old_name, new_name = clip_part.rsplit("/", 1)[1], emf_part.rsplit("/", 1)[1]
    emf = make_emf(make_pdf())
    out = zipfile.ZipFile(os.path.join(out_dir, "mac_clip.pptx"), "w", zipfile.ZIP_DEFLATED)
    for info in zin.infolist():
        data = zin.read(info.filename)
        name = info.filename
        if name == clip_part:
            name, data = emf_part, emf
        elif name.endswith(".rels"):
            data = data.replace(('/' + old_name + '"').encode(), ('/' + new_name + '"').encode())
        elif name == "[Content_Types].xml" and b'Extension="emf"' not in data:
            data = data.replace(b"<Default ", b'<Default Extension="emf" ContentType="image/x-emf"/><Default ', 1)
        out.writestr(name, data)
    out.close()
    print("fixture:", emf_part)


if __name__ == "__main__":
    build(sys.argv[1])
