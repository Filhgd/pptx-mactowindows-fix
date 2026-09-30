#!/usr/bin/env python3
"""Checks a fixed presentation against the original.

Usage: check_output.py ORIGINAL.pptx FIXED.pptx EXPECTED_WIDTH_PX"""
import io, re, sys, zipfile
import xml.etree.ElementTree as ET
from PIL import Image
from pptx import Presentation

orig_path, fixed_path, expected_w = sys.argv[1], sys.argv[2], int(sys.argv[3])
failures = []


def check(cond, msg):
    print(("OK   " if cond else "FAIL ") + msg)
    if not cond:
        failures.append(msg)


zo, zf = zipfile.ZipFile(orig_path), zipfile.ZipFile(fixed_path)
check(zf.testzip() is None, "zip is intact (all checksums)")
names = zf.namelist()
check(names[0] == "[Content_Types].xml", "[Content_Types].xml is the first part")

for n in names:
    if n.endswith(".xml") or n.endswith(".rels"):
        try:
            ET.fromstring(zf.read(n))
        except ET.ParseError as e:
            check(False, f"{n} is valid XML ({e})")
check(True, "all XML parts parse")

emfs = [n for n in zo.namelist() if n.endswith(".emf")]
for emf in emfs:
    check(emf not in names, f"{emf} removed")
    png = emf[:-4] + "_hr.png"
    check(png in names, f"{png} added")
    if png in names:
        im = Image.open(io.BytesIO(zf.read(png)))
        check(im.format == "PNG", "new image is a PNG")
        check(abs(im.size[0] - expected_w) <= 2, f"width {im.size[0]} px, expected {expected_w}")
        check(abs(im.size[0] / im.size[1] - 400 / 270) < 0.01, f"aspect ratio {im.size[0] / im.size[1]:.3f} (400/270)")
        rgba = im.convert("RGBA")
        alpha = rgba.getchannel("A")
        opaque = sum(alpha.histogram()[129:]) / (im.size[0] * im.size[1])
        check(0.01 < opaque < 0.6, f"drawn content covers {opaque:.1%} (transparent background, not empty)")
        # The blue square (PDF x 300-360, y 30-90 from the bottom) must be solid blue.
        sx = im.size[0] / 400
        px = rgba.getpixel((int(330 * sx), int((270 - 60) * sx)))
        check(px[2] > 200 and px[0] < 60 and px[3] > 250, f"blue square rendered at the right spot {px}")

allrels = "".join(zf.read(n).decode() for n in names if n.endswith(".rels"))
check(".emf" not in allrels, "no relationship points to an EMF any more")
ct = zf.read("[Content_Types].xml").decode()
check(re.search(r'Extension="png"', ct, re.I) is not None, "content type for PNG present")

changed = {n for n in names if n.endswith(".rels") or n == "[Content_Types].xml" or n.endswith("_hr.png")}
same = all(zo.read(n) == zf.read(n) for n in names if n not in changed)
check(same, "all other parts are byte-for-byte unchanged")

try:
    prs = Presentation(fixed_path)
    check(len(prs.slides) == len(Presentation(orig_path).slides), "python-pptx opens it, same number of slides")
except Exception as e:
    check(False, f"python-pptx opens it ({e})")

print("\n%d failure(s)" % len(failures))
sys.exit(1 if failures else 0)
